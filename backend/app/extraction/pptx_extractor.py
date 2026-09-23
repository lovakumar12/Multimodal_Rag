import io
from typing import List, Optional
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from backend.app.core.logging import logger
from backend.app.extraction.base import (
    BaseExtractor,
    ExtractedImageData,
    ExtractedPageData,
    ExtractedSectionData,
    ExtractedTableData,
    ExtractionResult,
)
from backend.app.extraction.vision_describer import vision_describer


class PPTXExtractor(BaseExtractor):
    """PowerPoint presentation extractor treating slides as pages and extracting text, tables, and images."""

    def extract(self, file_bytes: bytes, filename: str) -> ExtractionResult:
        prs = Presentation(io.BytesIO(file_bytes))
        pages_data: List[ExtractedPageData] = []
        slide_idx = 0

        for slide in prs.slides:
            slide_idx += 1
            page_num = slide_idx
            text_chunks: List[str] = []
            slide_title = None
            sections: List[ExtractedSectionData] = []
            tables: List[ExtractedTableData] = []
            images: List[ExtractedImageData] = []

            # Check for slide title
            if slide.shapes.title and slide.shapes.title.text:
                slide_title = slide.shapes.title.text.strip()
                sections.append(
                    ExtractedSectionData(
                        title=slide_title,
                        level=1,
                        page_number=page_num,
                        order_index=0,
                    )
                )

            # Traverse shapes
            for shape in slide.shapes:
                # Text Frame
                if shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        p_text = paragraph.text.strip()
                        if p_text and p_text != slide_title:
                            text_chunks.append(p_text)

                # Tables
                if shape.has_table:
                    tbl = shape.table
                    rows_data = []
                    for row in tbl.rows:
                        rows_data.append([cell.text.strip().replace("\n", " ") for cell in row.cells])
                    if rows_data and len(rows_data) >= 2:
                        headers = rows_data[0]
                        body_rows = rows_data[1:]
                        md_lines = [
                            "| " + " | ".join(headers) + " |",
                            "| " + " | ".join(["---"] * len(headers)) + " |",
                        ]
                        for r in body_rows:
                            padded = r + [""] * (len(headers) - len(r)) if len(r) < len(headers) else r[: len(headers)]
                            md_lines.append("| " + " | ".join(padded) + " |")
                        markdown = "\n".join(md_lines)
                        tables.append(
                            ExtractedTableData(
                                headers=headers,
                                rows=body_rows,
                                markdown=markdown,
                                page_number=page_num,
                                table_index=len(tables),
                                caption=f"Slide {page_num} Table",
                                surrounding_text=" ".join(text_chunks[:3]),
                            )
                        )

                # Pictures / Images
                if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    try:
                        image = shape.image
                        img_bytes = image.blob
                        img_ext = image.ext.upper() if image.ext else "PNG"

                        surrounding = f"Slide title: {slide_title or 'None'}. Text: {' '.join(text_chunks[:2])}"
                        desc, is_chart, caption = vision_describer.describe(
                            image_bytes=img_bytes,
                            image_format=img_ext,
                            surrounding_text=surrounding,
                            existing_caption=f"Slide {page_num} Image",
                        )

                        images.append(
                            ExtractedImageData(
                                image_bytes=img_bytes,
                                format=img_ext,
                                width=int(shape.width) if hasattr(shape, "width") else 0,
                                height=int(shape.height) if hasattr(shape, "height") else 0,
                                page_number=page_num,
                                caption=caption,
                                semantic_description=desc,
                                surrounding_text=surrounding,
                                is_diagram_or_chart=is_chart,
                            )
                        )
                    except Exception as e:
                        logger.warning(f"Error extracting image from slide {page_num}: {e}")

            # Speaker notes
            try:
                if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                    notes = slide.notes_slide.notes_text_frame.text.strip()
                    if notes:
                        text_chunks.append(f"[Speaker Notes]: {notes}")
            except Exception:
                pass

            full_text = "\n".join(text_chunks)
            if slide_title and not full_text.startswith(slide_title):
                full_text = f"# {slide_title}\n\n{full_text}"

            pages_data.append(
                ExtractedPageData(
                    page_number=page_num,
                    text=full_text,
                    sections=sections,
                    images=images,
                    tables=tables,
                )
            )

        return ExtractionResult(
            filename=filename,
            file_type="pptx",
            page_count=slide_idx,
            metadata={"title": filename},
            pages=pages_data,
        )
