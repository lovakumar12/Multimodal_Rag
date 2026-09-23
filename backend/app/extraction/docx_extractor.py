import io
from typing import List, Optional
import docx
from backend.app.core.logging import logger
from backend.app.extraction.base import (
    BaseExtractor,
    ExtractedImageData,
    ExtractedPageData,
    ExtractedSectionData,
    ExtractedTableData,
    ExtractionResult,
)
from backend.app.extraction.vision_describer import vision_describer


class DOCXExtractor(BaseExtractor):
    """DOCX extractor extracting headings, paragraphs, tables, and embedded images with logical pagination."""

    def extract(self, file_bytes: bytes, filename: str) -> ExtractionResult:
        doc = docx.Document(io.BytesIO(file_bytes))
        sections_data: List[ExtractedSectionData] = []
        tables_data: List[ExtractedTableData] = []
        images_data: List[ExtractedImageData] = []
        full_paragraphs: List[str] = []

        # 1. Extract paragraphs and headings
        for p in doc.paragraphs:
            text = p.text.strip()
            if not text:
                continue
            style_name = p.style.name.lower() if p.style and p.style.name else ""
            if "heading 1" in style_name:
                sections_data.append(ExtractedSectionData(title=text, level=1, order_index=len(sections_data)))
                full_paragraphs.append(f"\n# {text}\n")
            elif "heading 2" in style_name:
                sections_data.append(ExtractedSectionData(title=text, level=2, order_index=len(sections_data)))
                full_paragraphs.append(f"\n## {text}\n")
            elif "heading" in style_name:
                sections_data.append(ExtractedSectionData(title=text, level=3, order_index=len(sections_data)))
                full_paragraphs.append(f"\n### {text}\n")
            else:
                full_paragraphs.append(text)

        # 2. Extract tables
        for tbl_idx, tbl in enumerate(doc.tables):
            rows_data = []
            for row in tbl.rows:
                rows_data.append([cell.text.strip().replace("\n", " ") for cell in row.cells])
            if rows_data and len(rows_data) >= 2:
                headers = rows_data[0]
                body_rows = rows_data[1:]
                md_lines = [
                    "| " + " | ".join(headers) + " |",
                    "| " + " | ".join(["---"] * len(headers)) + " |",
                ]
                for r in body_rows:
                    padded = r + [""] * (len(headers) - len(r)) if len(r) < len(headers) else r[: len(headers)]
                    md_lines.append("| " + " | ".join(padded) + " |")
                markdown = "\n".join(md_lines)
                tables_data.append(
                    ExtractedTableData(
                        headers=headers,
                        rows=body_rows,
                        markdown=markdown,
                        page_number=1,
                        table_index=tbl_idx,
                        caption=f"Table {tbl_idx + 1}",
                        surrounding_text=" ".join(full_paragraphs[:3]),
                    )
                )

        # 3. Extract embedded images from document relationships
        for rel in doc.part.rels.values():
            if "image" in rel.target_ref:
                try:
                    img_part = rel.target_part
                    img_bytes = img_part.blob
                    ext = img_part.content_type.split("/")[-1].upper() if img_part.content_type else "PNG"
                    surrounding = " ".join(full_paragraphs[:3])
                    desc, is_chart, caption = vision_describer.describe(
                        image_bytes=img_bytes,
                        image_format=ext,
                        surrounding_text=surrounding,
                        existing_caption=f"DOCX Image {len(images_data) + 1}",
                    )
                    images_data.append(
                        ExtractedImageData(
                            image_bytes=img_bytes,
                            format=ext,
                            page_number=1,
                            caption=caption,
                            semantic_description=desc,
                            surrounding_text=surrounding,
                            is_diagram_or_chart=is_chart,
                        )
                    )
                except Exception as e:
                    logger.warning(f"Error extracting DOCX image: {e}")

        # Logical pagination (e.g. 2500 chars per page if long document)
        total_text = "\n\n".join(full_paragraphs)
        pages_data: List[ExtractedPageData] = []
        chunk_size = 2500
        if len(total_text) <= chunk_size:
            pages_data.append(
                ExtractedPageData(
                    page_number=1,
                    text=total_text,
                    sections=sections_data,
                    images=images_data,
                    tables=tables_data,
                )
            )
        else:
            slices = [total_text[i : i + chunk_size] for i in range(0, len(total_text), chunk_size)]
            for p_num, s in enumerate(slices, start=1):
                p_sections = [sec for sec in sections_data if sec.title in s]
                p_images = images_data if p_num == 1 else []
                p_tables = tables_data if p_num == 1 else []
                pages_data.append(
                    ExtractedPageData(
                        page_number=p_num,
                        text=s,
                        sections=p_sections,
                        images=p_images,
                        tables=p_tables,
                    )
                )

        return ExtractionResult(
            filename=filename,
            file_type="docx",
            page_count=len(pages_data),
            metadata={"title": filename},
            pages=pages_data,
        )
